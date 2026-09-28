"""Collect Machine Learning examples supported by boards in Simplicity Studio 6.

This script drives the already-open Simplicity Studio GUI through Microsoft UI
Automation.  It does not use hard-coded screen coordinates.

Edit BOARD_NAMES below, then run:
    python simplicity_studio_ml_inventory.py

Optional:
    python simplicity_studio_ml_inventory.py --boards BRD2608A BRD2608B BRD2605B
    python simplicity_studio_ml_inventory.py --output my_report.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pywinauto import Desktop
from pywinauto import mouse
from pywinauto.application import Application
from pywinauto.keyboard import send_keys
from pywinauto.timings import TimeoutError as PywinautoTimeoutError


# ---------------------------------------------------------------------------
# USER CONFIGURATION
# ---------------------------------------------------------------------------

BOARD_NAMES = [
    "BRD2608A",
    "BRD2608B",
    "BRD2605B",
]

MACHINE_LEARNING_FILTER = "machine learning"
OUTPUT_XLSX = "simplicity_studio_ml_support.xlsx"

STUDIO_CONNECT_TIMEOUT_SECONDS = 30
PAGE_LOAD_TIMEOUT_SECONDS = 45
SEARCH_TIMEOUT_SECONDS = 30
MAX_SCROLLS = 50
POLL_SECONDS = 0.35


class StudioAutomationError(RuntimeError):
    """Raised when the expected Studio state cannot be reached."""


@dataclass
class BoardResult:
    board_name: str
    studio_support: str
    supported_app_count: int
    supported_app_names: list[str]


def normalize(value: object) -> str:
    """Collapse line breaks/repeated spaces so UI names compare reliably."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def safe_name(control) -> str:
    try:
        return normalize(control.window_text())
    except Exception:
        return ""


def safe_control_type(control) -> str:
    try:
        return str(control.element_info.control_type or "")
    except Exception:
        return ""


def is_visible(control) -> bool:
    try:
        return bool(control.is_visible())
    except Exception:
        return False


def wait_until(predicate, timeout: float, description: str):
    deadline = time.monotonic() + timeout
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        try:
            value = predicate()
            if value:
                return value
        except Exception as exc:  # UI trees can refresh while being enumerated.
            last_error = exc
        time.sleep(POLL_SECONDS)
    suffix = f" Last UI error: {last_error}" if last_error else ""
    raise StudioAutomationError(f"Timed out waiting for {description}.{suffix}")


class StudioAutomation:
    WINDOW_TITLE_RE = re.compile(r"simplicity\s+studio", re.IGNORECASE)

    def __init__(self, connect_timeout: float = STUDIO_CONNECT_TIMEOUT_SECONDS):
        self.connect_timeout = connect_timeout
        self.app: Application | None = None
        self.window = None

    def connect(self) -> None:
        """Attach by enumerating top-level windows instead of matching one title."""
        def candidate_window():
            candidates = []
            for window in Desktop(backend="uia").windows():
                title = safe_name(window)
                if self.WINDOW_TITLE_RE.search(title):
                    try:
                        rect = window.rectangle()
                        area = max(0, rect.width()) * max(0, rect.height())
                    except Exception:
                        area = 0
                    candidates.append((is_visible(window), area, window))
            if not candidates:
                return None
            candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
            return candidates[0][2]

        desktop_window = wait_until(
            candidate_window,
            self.connect_timeout,
            "an open Simplicity Studio window",
        )
        handle = desktop_window.handle
        self.app = Application(backend="uia").connect(handle=handle, timeout=10)
        self.window = self.app.window(handle=handle)
        try:
            if self.window.is_minimized():
                self.window.restore()
        except Exception:
            pass
        self.window.set_focus()
        print(f"Connected to: {safe_name(self.window)}")

    def descendants(self, visible_only: bool = False) -> list:
        try:
            controls = self.window.descendants()
        except Exception:
            controls = []
        return [c for c in controls if not visible_only or is_visible(c)]

    def find_text_controls(
        self,
        text: str,
        *,
        exact: bool = True,
        visible_only: bool = True,
        control_types: Iterable[str] | None = None,
    ) -> list:
        expected = normalize(text).casefold()
        allowed = set(control_types or [])
        matches = []
        for control in self.descendants(visible_only=visible_only):
            if allowed and safe_control_type(control) not in allowed:
                continue
            actual = safe_name(control).casefold()
            if (actual == expected) if exact else (expected in actual):
                matches.append(control)
        return matches

    @staticmethod
    def click(control) -> None:
        try:
            control.scroll_into_view()
        except Exception:
            pass
        # Studio 289 is Chromium-based. UIA Invoke/Select can return without
        # activating some rendered tabs, while click_input() is the proven
        # behavior used by the working Studio_GUI_Automation_Pytest project.
        try:
            control.click_input()
            return
        except Exception:
            pass
        try:
            control.invoke()
            return
        except Exception:
            pass
        raise StudioAutomationError(
            f"Could not activate UI control {safe_name(control)!r}."
        )

    def click_named(
        self,
        text: str,
        *,
        exact: bool = True,
        timeout: float = PAGE_LOAD_TIMEOUT_SECONDS,
        control_types: Iterable[str] | None = None,
    ):
        def locate():
            matches = self.find_text_controls(
                text,
                exact=exact,
                visible_only=True,
                control_types=control_types,
            )
            if not matches:
                return None
            # Prefer actionable controls, then the smallest matching rectangle.
            priority = {
                "Button": 0,
                "TabItem": 1,
                "ListItem": 2,
                "DataItem": 3,
                "Hyperlink": 4,
                "Text": 5,
            }
            def rank(control):
                try:
                    rect = control.rectangle()
                    area = max(1, rect.width() * rect.height())
                except Exception:
                    area = 10**12
                return (priority.get(safe_control_type(control), 10), area)
            return sorted(matches, key=rank)[0]

        control = wait_until(locate, timeout, repr(text))
        self.click(control)
        return control

    def close_previous_tabs(self) -> None:
        """Close workspace tabs using UIA selection plus Ctrl+W, never coordinates."""
        for _ in range(20):
            tab_items = [
                c for c in self.descendants(visible_only=True)
                if safe_control_type(c) == "TabItem" and safe_name(c)
            ]
            # Board/project tabs are closed first. Home is harmless and reused.
            closable = [t for t in tab_items if safe_name(t).casefold() != "home"]
            if not closable:
                break
            tab = closable[-1]
            old_name = safe_name(tab)
            try:
                self.click(tab)
                send_keys("^w")
                wait_until(
                    lambda: not self.find_text_controls(
                        old_name,
                        exact=True,
                        visible_only=True,
                        control_types={"TabItem"},
                    ),
                    5,
                    f"tab {old_name!r} to close",
                )
            except Exception as exc:
                print(f"  Warning: could not close tab {old_name!r}: {exc}")
                break

    def open_device_catalog(self) -> None:
        self.click_named("DEVICES", exact=True, control_types={"Button", "Text", "Group"})
        self.click_named("Add Device(s)", exact=True, control_types={"Button", "Text"})
        wait_until(
            lambda: self.find_text_controls("Device & Products Catalog", exact=False),
            PAGE_LOAD_TIMEOUT_SECONDS,
            "Device & Products Catalog",
        )

    def visible_edits(self) -> list:
        # Never include Document here. Studio exposes its whole Chromium page as
        # a Document control; treating that as an input makes Ctrl+A select the
        # entire page (the blue screen shown in the failure screenshot).
        return [
            c for c in self.descendants(visible_only=True)
            if safe_control_type(c) == "Edit"
        ]

    @staticmethod
    def set_edit_text(edit, value: str, press_enter: bool = False) -> None:
        control_type = safe_control_type(edit)
        if control_type != "Edit":
            raise StudioAutomationError(
                f"Refusing to type into UIA control type {control_type!r}; "
                "expected a real Edit control."
            )
        try:
            edit.set_edit_text(value)
        except Exception:
            edit.click_input()
            try:
                if not edit.has_keyboard_focus():
                    raise StudioAutomationError(
                        "The search field did not receive keyboard focus."
                    )
            except AttributeError:
                # Older pywinauto builds may not expose has_keyboard_focus().
                pass
            send_keys("^a{BACKSPACE}")
            send_keys(value, with_spaces=True, pause=0.02)
        if press_enter:
            send_keys("{ENTER}")

    def catalog_search_edit(self):
        def locate():
            # Stable accessible name confirmed by the working automation.
            exact = [
                c for c in self.visible_edits()
                if safe_name(c).casefold()
                == "search by product name - kit, board or part"
            ]
            if exact:
                return exact[0]
            return None
        return wait_until(locate, PAGE_LOAD_TIMEOUT_SECONDS, "catalog search box")

    def catalog_result_count(self) -> int | None:
        pattern = re.compile(
            r"(\d+)\s+devices?(?:\(s\))?\s*(?:&|and)\s*"
            r"products?(?:\(s\))?",
            re.I,
        )
        for control in self.descendants(visible_only=True):
            match = pattern.search(safe_name(control))
            if match:
                return int(match.group(1))
        return None

    @staticmethod
    def catalog_has_no_results_message(names: Sequence[str]) -> bool:
        phrases = (
            "no devices",
            "no device",
            "no products",
            "no product",
            "no results",
            "no matching",
            "nothing found",
            "could not find",
            "0 devices",
            "0 device",
        )
        for name in names:
            folded = normalize(name).casefold()
            if any(phrase in folded for phrase in phrases):
                return True
        return False

    def search_board(self, board_name: str) -> bool:
        search = self.catalog_search_edit()
        self.window.set_focus()
        search.click_input()

        # Always clear the previous board search first. Studio can reuse the
        # same Home/catalog tab, so an old board ID may still be present.
        print(f"[{board_name}] Clearing previous board search...")
        search.type_keys("^a{BACKSPACE}", set_foreground=True)
        time.sleep(0.4)
        try:
            search.set_edit_text("")
        except Exception:
            pass

        print(f"[{board_name}] Entering board search...")
        search.click_input()
        search.type_keys(
            str(board_name),
            with_spaces=True,
            set_foreground=True,
        )
        time.sleep(0.5)

        query = normalize(board_name).casefold()
        search_started = time.monotonic()
        empty_state = {"signature": None, "stable_reads": 0}

        def resolved_search_state():
            controls = self.descendants(visible_only=True)
            names = [safe_name(control) for control in controls if safe_name(control)]

            # A real matching catalog row is the strongest success signal.
            matching_result = False
            for control in controls:
                ctype = safe_control_type(control)
                if ctype in {"Edit", "Document", "Window"}:
                    continue
                if query in safe_name(control).casefold():
                    matching_result = True
                    break
            if matching_result:
                return "SUPPORTED"

            count = self.catalog_result_count()
            if count == 0:
                return "UNSUPPORTED_NUMERIC_ZERO"

            if self.catalog_has_no_results_message(names):
                return "UNSUPPORTED_MESSAGE"

            loading = any(
                safe_control_type(control) == "ProgressBar"
                or "loading" in safe_name(control).casefold()
                or "please wait" in safe_name(control).casefold()
                for control in controls
            )
            if loading:
                empty_state["stable_reads"] = 0
                return False

            # Some Studio 289 builds remove the results table and its count
            # completely when there are no matches. In that state UIA exposes
            # neither "0 device(s)" nor a no-results label. Accept the empty
            # result only after the visible UI has remained stable for several
            # reads, so a catalog that is still rebuilding is not mislabeled.
            signature = tuple(sorted(set(names)))
            if signature == empty_state["signature"]:
                empty_state["stable_reads"] += 1
            else:
                empty_state["signature"] = signature
                empty_state["stable_reads"] = 1

            elapsed = time.monotonic() - search_started
            if elapsed >= 5.0 and empty_state["stable_reads"] >= 8:
                return "UNSUPPORTED_STABLE_EMPTY"
            return False

        state = wait_until(
            resolved_search_state,
            SEARCH_TIMEOUT_SECONDS,
            f"catalog results for {board_name}",
        )
        if state != "SUPPORTED":
            print(
                f"[{board_name}] Catalog search settled with no matching "
                f"board ({state})."
            )
            return False
        return True

    def select_board_result(self, board_name: str) -> None:
        query = normalize(board_name).casefold()

        def locate():
            candidates = []
            for control in self.descendants(visible_only=True):
                name = safe_name(control)
                if query not in name.casefold():
                    continue
                ctype = safe_control_type(control)
                if ctype in {"Edit", "Document", "Window"}:
                    continue
                score = {
                    "TreeItem": 0,
                    "Hyperlink": 1,
                    "Button": 2,
                    "ListItem": 3,
                    "DataItem": 4,
                    "Custom": 5,
                    "Text": 6,
                }.get(ctype, 100)
                candidates.append((score, len(name), control))
            if not candidates:
                return None
            candidates.sort(key=lambda item: (item[0], item[1]))
            return candidates[0][2]

        result = wait_until(locate, SEARCH_TIMEOUT_SECONDS, f"board result {board_name}")
        try:
            result.scroll_into_view()
            time.sleep(0.3)
        except Exception:
            pass
        # Resolve again after scrolling because Studio may rebuild the UI tree.
        result = wait_until(locate, SEARCH_TIMEOUT_SECONDS, f"board result {board_name}")
        result.click_input()
        wait_until(
            lambda: self.find_text_controls("EXAMPLE PROJECTS & DEMOS", exact=True),
            PAGE_LOAD_TIMEOUT_SECONDS,
            "board Overview page",
        )

    def example_filter_combo(self):
        def locate():
            for control in self.descendants(visible_only=True):
                if (
                    safe_control_type(control) == "ComboBox"
                    and safe_name(control).casefold() == "filter on keywords"
                ):
                    try:
                        if control.is_enabled():
                            return control
                    except Exception:
                        continue
            return None
        return wait_until(
            locate,
            PAGE_LOAD_TIMEOUT_SECONDS,
            "Filter on keywords ComboBox",
        )

    def wait_for_example_page_ready(self):
        def page_ready():
            combo = None
            for control in self.descendants(visible_only=True):
                name = safe_name(control).casefold()
                ctype = safe_control_type(control)
                if ctype == "ComboBox" and name == "filter on keywords":
                    try:
                        if control.is_enabled():
                            combo = control
                    except Exception:
                        pass
                if ctype == "ProgressBar" or "loading" in name or "please wait" in name:
                    return False
            return combo
        return wait_until(
            page_ready,
            PAGE_LOAD_TIMEOUT_SECONDS,
            "fully loaded Example Projects & Demos page",
        )

    def click_example_projects_and_demos(self) -> None:
        """Use the exact navigation approach from the working project."""
        def locate():
            candidates = []
            for control in self.descendants(visible_only=True):
                name = safe_name(control)
                normalized = " ".join(name.upper().split())
                if "EXAMPLE PROJECTS" not in normalized or "DEMOS" not in normalized:
                    continue
                try:
                    rect = control.rectangle()
                    if rect.width() <= 0 or rect.height() <= 0:
                        continue
                except Exception:
                    continue
                priority = {
                    "Hyperlink": 0,
                    "Button": 1,
                    "TabItem": 2,
                    "Text": 3,
                    "Custom": 4,
                }.get(safe_control_type(control), 100)
                candidates.append((priority, control))
            if not candidates:
                return None
            candidates.sort(key=lambda item: item[0])
            return candidates[0][1]

        target = wait_until(
            locate,
            PAGE_LOAD_TIMEOUT_SECONDS,
            "EXAMPLE PROJECTS & DEMOS navigation control",
        )
        print(
            "Using Example Projects control:",
            repr(safe_name(target)),
            "| Type:",
            safe_control_type(target),
        )
        self.window.set_focus()
        target.click_input()
        self.wait_for_example_page_ready()

    def open_examples_and_apply_filter(self) -> int:
        self.click_example_projects_and_demos()
        combo = self.example_filter_combo()
        self.window.set_focus()
        combo.click_input()
        time.sleep(0.3)
        combo.type_keys("^a{BACKSPACE}", set_foreground=True)
        time.sleep(0.5)
        combo.type_keys(
            MACHINE_LEARNING_FILTER,
            with_spaces=True,
            set_foreground=True,
        )
        time.sleep(0.3)
        combo.type_keys("{ENTER}", set_foreground=True)
        time.sleep(1.0)

        # Chromium can briefly show an intermediate count. Match the working
        # project and require three consecutive identical reads.
        stable = {"value": None, "observations": 0}
        def filtered_count_ready():
            count = self.items_found_count()
            if count is None:
                return False
            if count == stable["value"]:
                stable["observations"] += 1
            else:
                stable["value"] = count
                stable["observations"] = 1
            return count + 1 if stable["observations"] >= 3 else False

        count_plus_one = wait_until(
            filtered_count_ready,
            SEARCH_TIMEOUT_SECONDS,
            "filtered Items Found count",
        )
        return int(count_plus_one) - 1

    def items_found_count(self) -> int | None:
        pattern = re.compile(r"\b(\d+)\s+Items?\s+Found\b", re.I)
        for control in self.descendants(visible_only=True):
            match = pattern.search(safe_name(control))
            if match:
                return int(match.group(1))
        return None

    @staticmethod
    def plausible_title(text: str) -> bool:
        value = normalize(text)
        folded = value.casefold()
        rejected_exact = {
            "production", "create", "run demo", "more...", "more…",
            "view project documentation", "example projects", "demos",
            "solution examples", "filter on keywords", "machine learning",
        }
        rejected_starts = (
            "this ", "the ", "use ", "view ", "what are ",
            "wireless technology", "device type", "application type",
        )
        if folded in rejected_exact or folded.startswith(rejected_starts):
            return False
        if len(value) < 5 or len(value) > 180:
            return False
        if re.fullmatch(r"\d+\s+Items?\s+Found", value, re.I):
            return False
        return True

    def titles_from_visible_cards(self) -> set[str]:
        titles: set[str] = set()
        visible = self.descendants(visible_only=True)

        # Fast and reliable for Silicon Labs ML examples shown in Studio 289.
        for control in visible:
            name = safe_name(control)
            # The selected keyword chip is exposed as text named
            # "machine learning". It must never be counted as an app.
            if name.casefold() == MACHINE_LEARNING_FILTER.casefold():
                continue
            if re.match(r"^(AI\s*/\s*ML|MLTK)\b", name, re.I):
                if self.plausible_title(name):
                    titles.add(name)

        # Generic fallback: find the smallest card ancestor containing one CREATE.
        create_buttons = [
            c for c in visible
            if safe_name(c).casefold() == "create" and safe_control_type(c) == "Button"
        ]
        for button in create_buttons:
            ancestor = button
            card = None
            for _ in range(8):
                try:
                    ancestor = ancestor.parent()
                    create_count = sum(
                        1 for d in ancestor.descendants()
                        if safe_name(d).casefold() == "create"
                    )
                    if create_count == 1:
                        card = ancestor
                    elif create_count > 1:
                        break
                except Exception:
                    break
            if card is None:
                continue
            text_controls = []
            try:
                for child in card.descendants():
                    name = safe_name(child)
                    if safe_control_type(child) in {"Text", "Hyperlink"} and self.plausible_title(name):
                        rect = child.rectangle()
                        text_controls.append((rect.top, rect.left, name))
            except Exception:
                continue
            if text_controls:
                text_controls.sort()
                titles.add(text_controls[0][2])
        return titles

    def result_scroll_anchor(self):
        """Return a visible control located inside the application-card area."""
        anchors = []
        for control in self.descendants(visible_only=True):
            name = safe_name(control)
            ctype = safe_control_type(control)
            is_title = bool(
                re.match(r"^(AI\s*/\s*ML|MLTK)\b", name, re.I)
            )
            is_card_button = ctype == "Button" and name.casefold() in {
                "create",
                "run demo",
            }
            if not (is_title or is_card_button):
                continue
            try:
                rect = control.rectangle()
                if rect.width() > 0 and rect.height() > 0:
                    anchors.append((rect.top, rect.left, control))
            except Exception:
                pass
        if not anchors:
            return None
        # Prefer the lowest currently visible card so the pointer is certainly
        # over the vertically scrollable results region.
        anchors.sort(key=lambda item: (item[0], item[1]))
        return anchors[-1][2]

    def scroll_results_page(self, direction: str = "down") -> None:
        """Scroll the real card pane at a UIA-derived point, not a fixed point."""
        anchor = self.result_scroll_anchor()
        if anchor is None:
            raise StudioAutomationError(
                "Could not locate a visible application card to scroll."
            )
        rect = anchor.rectangle()
        point = (
            int((rect.left + rect.right) / 2),
            int((rect.top + rect.bottom) / 2),
        )
        mouse.scroll(
            coords=point,
            wheel_dist=-6 if direction == "down" else 6,
        )
        time.sleep(0.8)

    def collect_application_names(self, expected_count: int) -> list[str]:
        names: set[str] = set()
        unchanged = 0
        previous_size = -1

        if expected_count == 0:
            return []

        print("Reading visible application cards, then scrolling for more...")

        for scroll_number in range(MAX_SCROLLS + 1):
            visible_names = self.titles_from_visible_cards()
            new_names = sorted(visible_names - names, key=str.casefold)
            names.update(visible_names)
            for name in new_names:
                print(f"  Captured application: {name}")
            print(
                f"  Capture progress: {len(names)}/{expected_count} "
                f"after scroll step {scroll_number}"
            )
            if len(names) >= expected_count:
                break
            if len(names) == previous_size:
                unchanged += 1
            else:
                unchanged = 0
                previous_size = len(names)
            if unchanged >= 6:
                break
            self.scroll_results_page("down")

        result = sorted(names, key=str.casefold)
        if expected_count and len(result) != expected_count:
            raise StudioAutomationError(
                f"Studio reports {expected_count} filtered items, but {len(result)} unique "
                "application names were captured. Run Studio maximized and make sure the "
                "Examples page has finished loading."
            )
        return result

    def inspect_board(self, board_name: str) -> BoardResult:
        print(f"\n[{board_name}] Closing previous board tabs...")
        self.close_previous_tabs()
        print(f"[{board_name}] Opening Devices > Add Device(s)...")
        self.open_device_catalog()
        print(f"[{board_name}] Searching the device catalog...")
        if not self.search_board(board_name):
            print(f"[{board_name}] No catalog result: Studio Support = NO")
            return BoardResult(board_name, "NO", 0, [])

        print(f"[{board_name}] Selecting board and opening Example Projects & Demos...")
        self.select_board_result(board_name)
        count = self.open_examples_and_apply_filter()
        print(f"[{board_name}] Studio reports {count} Machine Learning item(s).")
        names = self.collect_application_names(count)
        for index, name in enumerate(names, start=1):
            print(f"  {index:02d}. {name}")
        return BoardResult(board_name, "YES", count, names)


def save_excel(results: Sequence[BoardResult], output_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Studio ML Support"

    headers = [
        "Sno",
        "Board Name",
        "Studio Support (Yes/No)",
        "Supported App Count",
        "Supported App Name",
    ]
    sheet.append(headers)
    for sno, result in enumerate(results, start=1):
        app_text = "\n".join(result.supported_app_names)
        if result.studio_support == "NO":
            app_text = "Not supported in Studio"
        sheet.append([
            sno,
            result.board_name,
            result.studio_support,
            result.supported_app_count,
            app_text,
        ])

    header_fill = PatternFill("solid", fgColor="0B7F83")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E1E8")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        status = row[2]
        if status.value == "YES":
            status.fill = PatternFill("solid", fgColor="E2F0D9")
            status.font = Font(color="276221", bold=True)
        else:
            status.fill = PatternFill("solid", fgColor="FCE4D6")
            status.font = Font(color="9C0006", bold=True)
        status.alignment = Alignment(horizontal="center", vertical="top")
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[3].alignment = Alignment(horizontal="center", vertical="top")

    widths = {"A": 8, "B": 22, "C": 24, "D": 22, "E": 58}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 32
    for row_number in range(2, sheet.max_row + 1):
        line_count = max(1, str(sheet.cell(row_number, 5).value or "").count("\n") + 1)
        sheet.row_dimensions[row_number].height = min(15 * line_count + 8, 300)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_suffix(".tmp.xlsx")
    workbook.save(temporary_path)
    temporary_path.replace(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inventory Machine Learning examples for boards in Simplicity Studio 6/289."
    )
    parser.add_argument(
        "--boards",
        nargs="+",
        help="Board IDs/names. If omitted, BOARD_NAMES inside this file is used.",
    )
    parser.add_argument(
        "--output",
        default=OUTPUT_XLSX,
        help=f"Excel report path (default: {OUTPUT_XLSX}).",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    boards = [normalize(item) for item in (args.boards or BOARD_NAMES) if normalize(item)]
    output_path = Path(args.output).expanduser().resolve()
    if not boards:
        print("ERROR: Add at least one board to BOARD_NAMES or pass --boards.", file=sys.stderr)
        return 2

    print("Simplicity Studio Machine Learning Inventory")
    print("Keep Simplicity Studio open, logged in, maximized, and unobstructed.")
    print(f"Boards: {', '.join(boards)}")
    print(f"Excel:  {output_path}")

    studio = StudioAutomation()
    results: list[BoardResult] = []
    try:
        studio.connect()
        for board in boards:
            try:
                result = studio.inspect_board(board)
            except KeyboardInterrupt:
                raise
            except Exception as exc:
                # A UI failure is different from a confirmed zero-result board.
                # Keep the partial report, then stop so it is never mislabeled NO.
                print(f"\nERROR while processing {board}: {exc}", file=sys.stderr)
                traceback.print_exc()
                save_excel(results, output_path)
                print(f"Partial Excel report saved: {output_path}")
                return 1
            results.append(result)
            save_excel(results, output_path)
            print(f"[{board}] Progress saved to {output_path.name}")
            if result.studio_support == "YES":
                # The selected board page is the active tab. Ctrl+W is a
                # coordinate-independent fallback if its close button is not
                # exposed as a UIA element.
                send_keys("^w")
                time.sleep(0.5)
    except KeyboardInterrupt:
        print("\nStopped by user. Saving collected results...")
        save_excel(results, output_path)
        return 130
    except (StudioAutomationError, PywinautoTimeoutError) as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        if results:
            save_excel(results, output_path)
        return 1
    finally:
        if studio.window is not None:
            try:
                studio.close_previous_tabs()
            except Exception:
                pass

    print(f"\nCompleted successfully. Excel report: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
